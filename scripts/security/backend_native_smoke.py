import gzip
import http.server
import json
import pathlib
import ssl
import subprocess
import tempfile
import threading
import wave

import fitz
import imageio_ffmpeg
import pdfplumber
import requests
import urllib3
from PIL import Image
from playwright.sync_api import sync_playwright

import importlib.util
assert importlib.util.find_spec('pip') is None, 'Build-only pip remains installed'
print('PASS: build-only pip removed')
status = pathlib.Path('/proc/self/status').read_text()
assert 'NoNewPrivs:	1' in status
assert 'CapEff:	0000000000000000' in status
print('PASS: privilege escalation disabled and effective capabilities empty')

assert urllib3.__version__ == '2.8.0', urllib3.__version__
print('urllib3:', urllib3.__version__)
print('requests:', requests.__version__)
# Dependency consistency was checked before pip removal during the image build.
print(subprocess.check_output(['dpkg-query', '-W', '-f=${Package} ${Version}\n',
    'libpcre2-8-0', 'libssl3t64', 'openssl', 'libxml2-16'], text=True))

class Handler(http.server.BaseHTTPRequestHandler):
    def do_GET(self):
        payload = gzip.compress(json.dumps({'ok': True}).encode())
        self.send_response(200)
        self.send_header('Content-Encoding', 'gzip')
        self.send_header('Transfer-Encoding', 'chunked')
        self.end_headers()
        self.wfile.write(f'{len(payload):X}\r\n'.encode() + payload + b'\r\n0\r\n\r\n')
    def log_message(self, *args):
        pass

server = http.server.ThreadingHTTPServer(('127.0.0.1', 0), Handler)
thread = threading.Thread(target=server.serve_forever, daemon=True)
thread.start()
try:
    with requests.Session() as session:
        session.trust_env = False
        response = session.get(f'http://127.0.0.1:{server.server_port}/', timeout=5)
        response.raise_for_status()
        assert response.json() == {'ok': True}
    print('PASS: requests/urllib3 chunked gzip response')
finally:
    server.shutdown()
    server.server_close()

with tempfile.TemporaryDirectory(prefix='lms-sbom-smoke-') as temp:
    root = pathlib.Path(temp)
    cert, key = root / 'cert.pem', root / 'key.pem'
    subprocess.run(['openssl', 'req', '-x509', '-newkey', 'rsa:2048', '-nodes',
        '-days', '1', '-subj', '/CN=localhost', '-addext',
        'subjectAltName=DNS:localhost,IP:127.0.0.1', '-keyout', str(key),
        '-out', str(cert)], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    tls_server = http.server.ThreadingHTTPServer(('127.0.0.1', 0), Handler)
    context = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
    context.load_cert_chain(cert, key)
    tls_server.socket = context.wrap_socket(tls_server.socket, server_side=True)
    threading.Thread(target=tls_server.serve_forever, daemon=True).start()
    try:
        with requests.Session() as session:
            session.trust_env = False
            url = f'https://127.0.0.1:{tls_server.server_port}/'
            assert session.get(url, verify=str(cert), timeout=5).json() == {'ok': True}
            try:
                session.get(url, timeout=5)
            except requests.exceptions.SSLError:
                pass
            else:
                raise AssertionError('Untrusted certificate was accepted')
        print('PASS: HTTPS trusted certificate accepted, untrusted certificate rejected')
    finally:
        tls_server.shutdown()
        tls_server.server_close()
    pdf = root / 'test.pdf'
    with fitz.open() as document:
        page = document.new_page()
        page.insert_text((72, 72), 'LMS dependency smoke test')
        document.save(pdf)
    with pdfplumber.open(pdf) as document:
        assert 'LMS dependency smoke test' in document.pages[0].extract_text()
    print('PASS: PDF generation and text extraction')

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        page = browser.new_page()
        page.set_content('<html><body><h1>LMS slide test</h1></body></html>')
        screenshot = root / 'slide.png'
        page.screenshot(path=str(screenshot))
        print('Chromium:', browser.version)
        browser.close()
    with Image.open(screenshot) as image:
        assert image.width > 0 and image.height > 0
    print('PASS: headless browser slide rendering')

    audio = root / 'audio.wav'
    with wave.open(str(audio), 'wb') as output:
        output.setnchannels(1)
        output.setsampwidth(2)
        output.setframerate(16000)
        output.writeframes(b'\0\0' * 16000)
    assert imageio_ffmpeg.get_ffmpeg_exe() == '/opt/ffmpeg/bin/ffmpeg'
    assert not list((pathlib.Path(imageio_ffmpeg.__file__).parent / 'binaries').glob('ffmpeg-*'))
    for name, executable in [('selected', imageio_ffmpeg.get_ffmpeg_exe())]:
        print(name, subprocess.check_output([executable, '-version'], text=True).splitlines()[0])
        video = root / f'{name}.mp4'
        subprocess.run([executable, '-hide_banner', '-loglevel', 'error', '-y',
            '-f', 'lavfi', '-i', 'color=c=blue:s=320x240:r=10', '-i', str(audio),
            '-t', '1', '-c:v', 'libx264', '-pix_fmt', 'yuv420p', '-c:a', 'aac',
            str(video)], check=True, timeout=30)
        assert video.stat().st_size > 0
        print(f'PASS: {name} FFmpeg video/audio encoding')
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    from app.documents.conversion import convert_office_to_pdf
    from app.generation.video import generate_hls_playlist
    from app.generation import tts
    document = Document()
    document.add_paragraph('Synthetic LMS DOCX conversion')
    document.add_picture(str(screenshot), width=Inches(4))
    docx_path = root / 'test.docx'
    document.save(docx_path)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1)).text = 'Synthetic LMS PPTX conversion'
    slide.shapes.add_picture(str(screenshot), Inches(1), Inches(3), width=Inches(4))
    pptx_path = root / 'test.pptx'
    presentation.save(pptx_path)
    for office_path in (docx_path, pptx_path):
        pdf_path = convert_office_to_pdf(office_path, office_path.with_suffix('.pdf'))
        with fitz.open(pdf_path) as converted:
            assert 'Synthetic LMS' in ''.join(page.get_text() for page in converted)
            assert any(page.get_images() for page in converted), 'Embedded image lost during conversion'
        print('PASS: actual LMS Office-to-PDF conversion:', office_path.suffix)
    playlist = generate_hls_playlist(video)
    assert '#EXTM3U' in playlist.read_text()
    assert (playlist.parent / 'init.mp4').stat().st_size > 0
    assert list(playlist.parent.glob('segment_*.m4s'))
    print('PASS: actual LMS fMP4 HLS packaging')
    previous_speed = tts.TTS_SPEED
    try:
        tts.TTS_SPEED = 1.2
        assert tts._apply_tts_speed(str(audio))
        with wave.open(str(audio)) as output:
            assert output.getnframes() > 0
        print('PASS: actual LMS narration speed adjustment')
    finally:
        tts.TTS_SPEED = previous_speed
from lxml import etree
assert etree.LIBXML_VERSION >= (2, 15, 4), etree.LIBXML_VERSION
print('Wheel libxml2:', etree.LIBXML_VERSION)
for module in pathlib.Path('/usr/local/lib/python3.12/lib-dynload').glob('*.so'):
    linked = subprocess.run(['ldd', str(module)], capture_output=True, text=True, check=False)
    assert 'not found' not in linked.stdout + linked.stderr, (str(module), linked.stdout)
print('PASS: every CPython native extension has resolved runtime libraries')
from app.generation.video import capture_slide_frames, encode_slide_clip, concatenate_clips
from app.generation import video as video_module
original_inject_assets = video_module._inject_local_slide_assets
def inject_assets_with_test_fonts(page):
    # The container is offline: stub only the external font stylesheet, not LMS assets.
    page.route('https://fonts.googleapis.com/**', lambda route: route.fulfill(status=200, content_type='text/css', body='/* synthetic test font response */'))
    original_inject_assets(page)
video_module._inject_local_slide_assets = inject_assets_with_test_fonts
with tempfile.TemporaryDirectory(prefix='lms-real-media-') as tmp:
    root = pathlib.Path(tmp)
    html = root / 'slides.html'
    html.write_text('<html><body><div class="slide" id="slide">Slide one</div><script>window.goToSlide = i => document.getElementById("slide").textContent = "Synthetic slide " + i;</script></body></html>')
    frames = capture_slide_frames(str(html), slide_count=2, output_dir=str(root))
    audio = root / 'narration.wav'
    with wave.open(str(audio), 'wb') as output:
        output.setnchannels(1); output.setsampwidth(2); output.setframerate(16000); output.writeframes(b'\0\0' * 16000)
    clips=[]
    for i, frame in enumerate(frames):
        clip = root / f'clip-{i}.mp4'
        encode_slide_clip(frame, str(audio), str(clip), transition_pause_seconds=0.1)
        assert clip.stat().st_size > 0
        clips.append(str(clip))
    combined = root / 'combined.mp4'
    concatenate_clips(clips, str(combined), working_dir=str(root))
    playlist = generate_hls_playlist(combined)
    assert playlist.is_file()
    print('PASS: actual LMS slide capture, narration encoding, clip concatenation and HLS')
print('All candidate dependency smoke checks passed.')
