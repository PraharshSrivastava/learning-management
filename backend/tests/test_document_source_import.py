from io import BytesIO

import fitz
import pytest
from docx import Document
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.core.settings import settings
from app.document_builder import import_source
from app.schemas.document_workflow import DraftContent, ImageBlock, ImportRequest, TableBlock

pytest_plugins = ["test_document_draft_repository"]


def source(repo, tmp_path, monkeypatch, suffix, build):
    monkeypatch.setattr(settings, "upload_dir", tmp_path / "uploads")
    monkeypatch.setattr(settings, "derived_document_dir", tmp_path / "derived")
    settings.upload_dir.mkdir()
    path = settings.upload_dir / ("sample" + suffix)
    build(path)
    monkeypatch.setattr(
        import_source,
        "get_document",
        lambda _: {
            "trainer_id": "owner",
            "file_path": path.name,
            "file_name": path.name,
        },
    )
    draft = repo.create("owner", DraftContent(entryMode="existing"))
    return import_source.SourceImporter(repo).load(
        draft.id, "owner", ImportRequest(expectedRevision=1, documentId="synthetic")
    )


def png():
    buffer = BytesIO()
    Image.new("RGB", (30, 20), "blue").save(buffer, "PNG")
    return buffer.getvalue()


def test_docx_preserves_module_order_table_and_owned_image(repo, tmp_path, monkeypatch):
    def build(path):
        document = Document()
        document.add_heading("Module 1: Listening", 1)
        document.add_paragraph("Listen before proposing a solution.")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Action"
        table.cell(0, 1).text = "Purpose"
        table.cell(1, 0).text = "Clarify"
        table.cell(1, 1).text = "Understand the concern"
        document.add_picture(BytesIO(png()))
        document.add_heading("Module 2: Follow-up", 1)
        document.add_paragraph("Confirm the next step.")
        document.save(path)

    draft = source(repo, tmp_path, monkeypatch, ".docx", build)
    assert [module.title for module in draft.content.modules] == ["Listening", "Follow-up"]
    assert any(isinstance(block, TableBlock) for block in draft.content.modules[0].blocks)
    assert any(isinstance(block, ImageBlock) for block in draft.content.modules[0].blocks)
    resources = repo.resources(draft.id, "owner")
    assert len([r for r in resources if r["kind"] == "asset"]) == 1
    repo.validate_resources(draft.id, "owner", draft.content)


def test_pdf_needs_no_lms_metadata_table_and_reports_unreadable_pages(repo, tmp_path, monkeypatch):
    def build(path):
        with fitz.open() as document:
            page = document.new_page()
            page.insert_text((50, 50), "Module 1: Customer needs\nListen first.")
            document.new_page()
            document.save(path)

    draft = source(repo, tmp_path, monkeypatch, ".pdf", build)
    assert draft.content.modules[0].title == "Customer needs"
    record = repo.resources(draft.id, "owner", "source")[0]
    assert any("Page 2" in issue for issue in record["issues"])
    assert record["acknowledged"] is False


def test_pptx_preserves_slide_text_and_embedded_image(repo, tmp_path, monkeypatch):
    def build(path):
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[5])
        slide.shapes.title.text = "Customer needs"
        text = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        text.text = "Confirm your understanding."
        slide.shapes.add_picture(BytesIO(png()), Inches(1), Inches(3))
        document.save(path)

    draft = source(repo, tmp_path, monkeypatch, ".pptx", build)
    assert "Confirm your understanding." in "\n".join(
        getattr(b, "text", "") for b in draft.content.modules[0].blocks
    )
    assert any(isinstance(b, ImageBlock) for b in draft.content.modules[0].blocks)


def test_reference_attachment_does_not_replace_existing_document(repo, tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "derived_document_dir", tmp_path / "derived")
    draft = repo.create("owner", DraftContent(entryMode="ai"))
    attached = import_source.SourceImporter(repo).load(
        draft.id,
        "owner",
        ImportRequest(expectedRevision=1, referenceOnly=True, text="Listen without interrupting."),
    )
    assert attached.content == draft.content
    assert len(repo.resources(draft.id, "owner", "source")) == 1


def test_image_only_pdf_is_rejected_without_silent_empty_success(repo, tmp_path, monkeypatch):
    def build(path):
        with fitz.open() as document:
            document.new_page().insert_image(fitz.Rect(10, 10, 100, 100), stream=png())
            document.save(path)

    with pytest.raises(ValueError, match="No readable text"):
        source(repo, tmp_path, monkeypatch, ".pdf", build)
